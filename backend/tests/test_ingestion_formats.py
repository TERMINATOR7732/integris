"""Comprehensive tests for multi-format ingestion (CSV, TSV, TXT, XLSX, XLS, PDF)."""

import io
from pathlib import Path
from fastapi.testclient import TestClient
import openpyxl
import pandas as pd
import pytest

from app.ingestion import (
    detect_file_type,
    ingest_dataset,
    parse_csv,
    parse_excel,
    parse_pdf,
    parse_txt,
)
from app.main import app
from app.models.report import ForensicDossier

client = TestClient(app)

# Helper to build a minimal valid PDF containing an explicit tabular grid
def create_minimal_pdf_with_table() -> bytes:
    stream_content = (
        b"BT\n"
        b"/F1 10 Tf\n"
        b"1 0 0 1 110 685 Tm (employee_id) Tj\n"
        b"1 0 0 1 210 685 Tm (department) Tj\n"
        b"1 0 0 1 110 665 Tm (EMP-101) Tj\n"
        b"1 0 0 1 210 665 Tm (Engineering) Tj\n"
        b"1 0 0 1 110 645 Tm (EMP-102) Tj\n"
        b"1 0 0 1 210 645 Tm (Finance) Tj\n"
        b"ET\n"
        b"100 700 m 300 700 l S\n"
        b"100 680 m 300 680 l S\n"
        b"100 660 m 300 660 l S\n"
        b"100 640 m 300 640 l S\n"
        b"100 700 m 100 640 l S\n"
        b"200 700 m 200 640 l S\n"
        b"300 700 m 300 640 l S\n"
    )
    pdf_template = (
        b"%%PDF-1.4\n"
        b"1 0 obj <</Type /Catalog /Pages 2 0 R>> endobj\n"
        b"2 0 obj <</Type /Pages /Kids [3 0 R] /Count 1>> endobj\n"
        b"3 0 obj <</Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources 4 0 R /Contents 5 0 R>> endobj\n"
        b"4 0 obj <</Font <</F1 <</Type /Font /Subtype /Type1 /BaseFont /Helvetica>>>>>> endobj\n"
        b"5 0 obj <</Length %d>> stream\n%sendstream\nendobj\n"
        b"xref\n0 6\n"
        b"0000000000 65535 f \n"
        b"0000000010 00000 n \n"
        b"0000000060 00000 n \n"
        b"0000000117 00000 n \n"
        b"0000000216 00000 n \n"
        b"0000000293 00000 n \n"
        b"trailer <</Size 6 /Root 1 0 R>>\n"
        b"startxref\n650\n%%%%EOF\n"
    ) % (len(stream_content), stream_content)
    return pdf_template


# ==============================================================================
# 1. Format Detection & Magic Bytes Tests
# ==============================================================================

def test_detect_file_type_csv():
    assert detect_file_type("test.csv", b"id,val\n1,2") == "csv"
    assert detect_file_type("test.tsv", b"id\tval\n1\t2") == "tsv"
    assert detect_file_type("test.txt", b"id,val\n1,2") == "txt"


def test_detect_file_type_xlsx():
    zip_header = b"PK\x03\x04\x14\x00\x00\x00"
    assert detect_file_type("data.xlsx", zip_header) == "xlsx"


def test_detect_file_type_pdf():
    pdf_header = b"%PDF-1.7\n..."
    assert detect_file_type("report.pdf", pdf_header) == "pdf"


def test_detect_file_type_mismatch():
    with pytest.raises(ValueError, match="PDF"):
        detect_file_type("report.csv", b"%PDF-1.4...")

    with pytest.raises(ValueError, match="OpenXML"):
        detect_file_type("report.csv", b"PK\x03\x04...")

    with pytest.raises(ValueError, match="Unsupported"):
        detect_file_type("program.exe", b"MZ...")


# ==============================================================================
# 2. Delimited CSV/TSV Ingestion Tests
# ==============================================================================

def test_parse_csv_delimiters():
    # Comma
    df_comma = parse_csv(b"col1,col2,col3\n1,2,3\n4,5,6")
    assert list(df_comma.columns) == ["col1", "col2", "col3"]
    assert len(df_comma) == 2

    # Tab
    df_tab = parse_csv(b"col1\tcol2\tcol3\n1\t2\t3\n4\t5\t6")
    assert list(df_tab.columns) == ["col1", "col2", "col3"]
    assert len(df_tab) == 2

    # Semicolon
    df_semi = parse_csv(b"col1;col2;col3\n1;2;3\n4;5;6")
    assert list(df_semi.columns) == ["col1", "col2", "col3"]
    assert len(df_semi) == 2

    # Pipe
    df_pipe = parse_csv(b"col1|col2|col3\n1|2|3\n4|5|6")
    assert list(df_pipe.columns) == ["col1", "col2", "col3"]
    assert len(df_pipe) == 2


# ==============================================================================
# 3. TXT Format Ingestion & Prose Rejection Tests
# ==============================================================================

def test_parse_txt_valid_tabular():
    txt_content = b"user_id,activity,score\nU101,login,95.5\nU102,logout,88.0\nU103,update,72.4"
    df = parse_txt(txt_content)
    assert len(df) == 3
    assert list(df.columns) == ["user_id", "activity", "score"]


def test_parse_txt_rejects_unstructured_prose():
    prose = (
        b"The quick brown fox jumps over the lazy dog.\n"
        b"Integris is an advanced data forensics platform.\n"
        b"Zero retention volatile memory is guaranteed by design.\n"
        b"This file is pure natural language prose without tabular structure.\n"
    )
    with pytest.raises(ValueError, match="structured tabular data"):
        parse_txt(prose)


def test_parse_txt_empty_raises():
    with pytest.raises(ValueError, match="empty"):
        parse_txt(b"   \n  \t  \n  ")


# ==============================================================================
# 4. Excel (XLSX) Ingestion Tests
# ==============================================================================

def test_parse_excel_single_sheet():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Employees"
    ws.append(["id", "name", "salary"])
    ws.append([1, "Alice", 95000])
    ws.append([2, "Bob", 82000])

    bio = io.BytesIO()
    wb.save(bio)
    xlsx_bytes = bio.getvalue()

    df, sheet, sheets = parse_excel(xlsx_bytes, "xlsx")
    assert sheet == "Employees"
    assert sheets == ["Employees"]
    assert len(df) == 2
    assert list(df.columns) == ["id", "name", "salary"]


def test_parse_excel_multi_sheet_skips_empty():
    wb = openpyxl.Workbook()
    # Sheet 1: Empty
    ws1 = wb.active
    ws1.title = "CoverPage"
    # Sheet 2: Tabular data
    ws2 = wb.create_sheet(title="Financials")
    ws2.append(["quarter", "revenue", "profit"])
    ws2.append(["Q1", 100000, 25000])
    ws2.append(["Q2", 120000, 30000])

    bio = io.BytesIO()
    wb.save(bio)
    xlsx_bytes = bio.getvalue()

    df, sheet, sheets = parse_excel(xlsx_bytes, "xlsx")
    assert sheet == "Financials"
    assert "CoverPage" in sheets
    assert "Financials" in sheets
    assert len(df) == 2
    assert list(df.columns) == ["quarter", "revenue", "profit"]


def test_parse_excel_all_empty_sheets_raises():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "EmptySheet"

    bio = io.BytesIO()
    wb.save(bio)
    xlsx_bytes = bio.getvalue()

    with pytest.raises(ValueError, match="no tabular records"):
        parse_excel(xlsx_bytes, "xlsx")


# ==============================================================================
# 5. PDF Ingestion Tests
# ==============================================================================

def test_parse_pdf_valid_table():
    pdf_bytes = create_minimal_pdf_with_table()
    df, page_count, table_index = parse_pdf(pdf_bytes)
    assert page_count == 1
    assert table_index == 1
    assert list(df.columns) == ["employee_id", "department"]
    assert len(df) == 2
    assert df.iloc[0]["employee_id"] == "EMP-101"
    assert df.iloc[1]["department"] == "Finance"


def test_parse_pdf_non_tabular_certificate_rejected():
    cert_path = Path(__file__).resolve().parent.parent.parent / "Manas_Anil_Kulkarni_3412_Certificate.pdf"
    if cert_path.exists():
        with open(cert_path, "rb") as f:
            cert_bytes = f.read()
        with pytest.raises(ValueError, match="machine-readable tabular data"):
            parse_pdf(cert_bytes)


def test_parse_pdf_numeric_inference_and_text_preservation(monkeypatch):
    """Verify PDF table extraction infers numeric columns while preserving IDs, dates, text, and mixed columns."""
    from app.engine.profiler import profile_dataset
    from app.models.report import SemanticType

    class DummyPage:
        def extract_tables(self):
            return [[
                ["employee_id", "user_id", "hire_date", "department", "amount", "count", "mixed_col"],
                ["EMP-101", "1001", "2026-09-25", "Engineering", "100", "10", "100"],
                ["EMP-102", "1002", "2026-09-26", "Finance", "42.5", "-17", "unknown"],
                ["EMP-103", "1003", "2026-09-27", "Legal", "-17", "25", "42.5"],
                ["EMP-104", "1004", "2026-09-28", "Operations", "88.25", "0", "99"],
                ["EMP-105", "1005", "2026-09-29", "Engineering", "150", "5", "85"],
            ]]

    class DummyPDF:
        def __init__(self):
            self.pages = [DummyPage()]
        def __enter__(self):
            return self
        def __exit__(self, *args):
            pass

    import pdfplumber
    monkeypatch.setattr(pdfplumber, "open", lambda _: DummyPDF())

    df, page_count, _ = parse_pdf(b"%PDF-1.4 dummy")
    assert page_count == 1

    # Numeric columns must have numeric dtypes
    assert pd.api.types.is_float_dtype(df["amount"])
    assert pd.api.types.is_integer_dtype(df["count"])
    assert df["amount"].tolist() == [100.0, 42.5, -17.0, 88.25, 150.0]
    assert df["count"].tolist() == [10, -17, 25, 0, 5]

    # Identifiers, dates, text, and mixed columns must remain string/object (not coerced to numeric)
    for str_col in ("employee_id", "user_id", "hire_date", "department", "mixed_col"):
        assert pd.api.types.is_string_dtype(df[str_col]) or df[str_col].dtype == object
        assert not pd.api.types.is_numeric_dtype(df[str_col])

    # Verify downstream profiler receives numeric dtypes and computes stats
    _, profiles = profile_dataset(df)
    prof_map = {p.name: p for p in profiles}
    assert prof_map["amount"].semantic_type == SemanticType.NUMERIC_CONTINUOUS
    assert prof_map["amount"].min_value == -17.0
    assert prof_map["amount"].max_value == 150.0


def test_parse_pdf_multipage_coalesces_only_valid_tables(monkeypatch):
    """Verify valid multi-page tables are coalesced while rejected tables in extracted_tables do not contaminate rows."""
    rejected_banner_table = [
        [" Decorative Header Box With Multi-line Prose\nLine 2\nLine 3\nLine 4 ", " Banner Side\nLine 2\nLine 3\nLine 4 "],
        [" Paragraph block one\nLine 2\nLine 3\nLine 4 ", " Paragraph block two\nLine 2\nLine 3\nLine 4 "],
    ]
    valid_page1_table = [
        ["employee_id", "score"],
        ["EMP-01", "100"],
        ["EMP-02", "42.5"],
    ]
    valid_page2_table = [
        ["employee_id", "score"],
        ["EMP-03", "-17"],
        ["EMP-04", "85"],
    ]
    rejected_footer_table = [
        [" Footer disclaimer\nLine 2\nLine 3\nLine 4 ", " Signature block\nLine 2\nLine 3\nLine 4 "],
        [" Legal text\nLine 2\nLine 3\nLine 4 ", " Stamp area\nLine 2\nLine 3\nLine 4 "],
    ]

    class DummyPage:
        def __init__(self, tables):
            self._tables = tables
        def extract_tables(self):
            return self._tables

    class DummyPDF:
        def __init__(self):
            self.pages = [
                DummyPage([rejected_banner_table, valid_page1_table]),
                DummyPage([valid_page2_table, rejected_footer_table]),
            ]
        def __enter__(self):
            return self
        def __exit__(self, *args):
            pass

    import pdfplumber
    monkeypatch.setattr(pdfplumber, "open", lambda _: DummyPDF())

    df, page_count, _ = parse_pdf(b"%PDF-1.4 dummy")
    assert page_count == 2
    assert list(df.columns) == ["employee_id", "score"]
    assert len(df) == 4
    assert df["employee_id"].tolist() == ["EMP-01", "EMP-02", "EMP-03", "EMP-04"]
    assert df["score"].tolist() == [100.0, 42.5, -17.0, 85.0]



# ==============================================================================
# 6. End-to-End API Integration Tests (POST /api/v1/investigate)
# ==============================================================================

def test_api_investigate_xlsx():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Forensics"
    ws.append(["record_id", "feature_a", "feature_b"])
    for i in range(10):
        ws.append([f"REC-{i}", i * 10, i * 2.5])

    bio = io.BytesIO()
    wb.save(bio)

    response = client.post(
        "/api/v1/investigate",
        files={"file": ("dataset.xlsx", bio.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert response.status_code == 200
    data = response.json()
    dossier = ForensicDossier.model_validate(data)
    assert dossier.metadata.file_name == "dataset.xlsx"
    assert dossier.metadata.file_type == "xlsx"
    assert dossier.metadata.sheet_name == "Forensics"
    assert dossier.metadata.row_count == 10
    assert dossier.metadata.column_count == 3


def test_api_investigate_pdf():
    pdf_bytes = create_minimal_pdf_with_table()
    response = client.post(
        "/api/v1/investigate",
        files={"file": ("report.pdf", pdf_bytes, "application/pdf")},
    )
    assert response.status_code == 200
    data = response.json()
    dossier = ForensicDossier.model_validate(data)
    assert dossier.metadata.file_name == "report.pdf"
    assert dossier.metadata.file_type == "pdf"
    assert dossier.metadata.page_count == 1
    assert dossier.metadata.row_count == 2
    assert dossier.metadata.column_count == 2


def test_api_investigate_txt_tabular():
    txt_data = b"id,category,score\n1,A,90\n2,B,85\n3,C,92"
    response = client.post(
        "/api/v1/investigate",
        files={"file": ("data.txt", txt_data, "text/plain")},
    )
    assert response.status_code == 200
    data = response.json()
    dossier = ForensicDossier.model_validate(data)
    assert dossier.metadata.file_name == "data.txt"
    assert dossier.metadata.file_type == "txt"
    assert dossier.metadata.row_count == 3


def test_api_investigate_txt_prose_rejected():
    prose_data = b"This is a prose article explaining the concept of zero-retention architecture in modern cloud applications."
    response = client.post(
        "/api/v1/investigate",
        files={"file": ("article.txt", prose_data, "text/plain")},
    )
    assert response.status_code == 400
    assert "structured tabular data" in response.json()["detail"].lower()


def test_api_investigate_size_limit_excel():
    # 26 MB dummy payload with .xlsx extension
    large_payload = b"PK\x03\x04" + (b"0" * (26 * 1024 * 1024))
    response = client.post(
        "/api/v1/investigate",
        files={"file": ("huge.xlsx", large_payload, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert response.status_code == 413
    assert "exceeds the maximum allowed limit of 25 mb" in response.json()["detail"].lower()


def test_api_investigate_size_limit_pdf():
    # 16 MB dummy payload with .pdf extension
    large_payload = b"%PDF-1.4" + (b"0" * (16 * 1024 * 1024))
    response = client.post(
        "/api/v1/investigate",
        files={"file": ("huge.pdf", large_payload, "application/pdf")},
    )
    assert response.status_code == 413
    assert "exceeds the maximum allowed limit of 15 mb" in response.json()["detail"].lower()


def test_parse_csv_and_txt_utf8_bom():
    """Ensure UTF-8 BOM is cleanly stripped from the first column header in CSV and TXT."""
    bom_csv = b"\xef\xbb\xbfemployee_id,department,salary\nE1,Sales,50000\nE2,IT,60000"
    df_csv = parse_csv(bom_csv)
    assert list(df_csv.columns) == ["employee_id", "department", "salary"]

    df_txt = parse_txt(bom_csv)
    assert list(df_txt.columns) == ["employee_id", "department", "salary"]


def test_detect_file_type_rejects_binary_null_bytes_in_text():
    """Ensure binary streams containing null bytes cannot masquerade as .csv, .tsv, or .txt."""
    binary_payload = b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00"
    with pytest.raises(ValueError, match="binary null-byte"):
        detect_file_type("fake.csv", binary_payload)
    with pytest.raises(ValueError, match="binary null-byte"):
        detect_file_type("fake.txt", binary_payload)


def test_parse_excel_zip_bomb_and_duplicate_headers(monkeypatch):
    """Verify XLSX zip-bomb uncompressed size limit and duplicate header deduplication after blank rows."""
    import app.ingestion.excel_parser as excel_mod

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Audit"
    ws.append(["user_id", "score", "score "])  # Duplicate header after .strip()
    ws.append([None, None, None])  # Blank row between header and data
    ws.append(["U1", 10, 20])
    ws.append(["U1", 10, 20])  # Exact duplicate row after dropped blank row

    bio = io.BytesIO()
    wb.save(bio)
    xlsx_bytes = bio.getvalue()

    df, sheet, _ = parse_excel(xlsx_bytes, "xlsx")
    assert sheet == "Audit"
    assert list(df.columns) == ["user_id", "score", "score_1"]
    assert list(df.index) == [0, 1]

    # Verify zip-bomb limit triggers before openpyxl parses
    monkeypatch.setattr(excel_mod, "MAX_EXCEL_UNCOMPRESSED_BYTES", 100)
    with pytest.raises(ValueError, match="decompression bomb"):
        parse_excel(xlsx_bytes, "xlsx")


def test_parse_pdf_incompatible_page_schema_skipped(monkeypatch):
    """Ensure PDF pages with incompatible column counts do not corrupt the primary table."""
    page1_table = [
        ["employee_id", "department"],
        ["EMP-01", "Engineering"],
        ["EMP-02", "Finance"],
    ]
    page2_incompatible_table = [
        ["summary_metric", "q1", "q2", "q3"],
        ["Revenue", "100", "200", "300"],
        ["Cost", "50", "60", "70"],
    ]

    class DummyPage:
        def __init__(self, tables):
            self._tables = tables
        def extract_tables(self):
            return self._tables

    class DummyPDF:
        def __init__(self):
            self.pages = [DummyPage([page1_table]), DummyPage([page2_incompatible_table])]
        def __enter__(self):
            return self
        def __exit__(self, *args):
            pass

    import pdfplumber
    monkeypatch.setattr(pdfplumber, "open", lambda _: DummyPDF())

    df, page_count, _ = parse_pdf(b"%PDF-1.4 dummy")
    assert page_count == 2
    assert list(df.columns) == ["employee_id", "department"]
    assert len(df) == 2


def test_detect_file_type_pdf_with_excel_magic_bytes_rejected():
    """Ensure a .pdf filename containing ZIP/XLSX or OLE2/XLS magic bytes is rejected."""
    with pytest.raises(ValueError, match="OpenXML/ZIP"):
        detect_file_type("disguised.pdf", b"PK\x03\x04\x14\x00\x00\x00")

    with pytest.raises(ValueError, match="legacy OLE/XLS"):
        detect_file_type("disguised.pdf", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1\x00\x00")


def test_api_corrupted_excel_does_not_leak_internal_exceptions():
    """Ensure corrupted .xlsx and .xls uploads return clean 400 messages without internal library tracebacks."""
    corrupted_xlsx = b"PK\x03\x04corrupted_not_a_valid_zip_archive_directory"
    resp_xlsx = client.post(
        "/api/v1/investigate",
        files={"file": ("broken.xlsx", corrupted_xlsx, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert resp_xlsx.status_code == 400
    detail_xlsx = resp_xlsx.json()["detail"]
    assert "Unable to read Excel file (.xlsx)" in detail_xlsx
    assert "BadZipFile" not in detail_xlsx
    assert "File is not a zip file" not in detail_xlsx

    corrupted_xls = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1corrupted_ole2_stream_payload"
    resp_xls = client.post(
        "/api/v1/investigate",
        files={"file": ("broken.xls", corrupted_xls, "application/vnd.ms-excel")},
    )
    assert resp_xls.status_code == 400
    detail_xls = resp_xls.json()["detail"]
    assert "Unable to read Excel file (.xls)" in detail_xls
    assert "XLRDError" not in detail_xls


def test_parse_excel_preserves_na_text_sentinels_and_blank_cells():
    """Verify Excel ingestion preserves literal 'N/A' strings for type-drift/sentinel detection while treating blank cells as NaN."""
    from app.engine.completeness import analyze_completeness
    from app.engine.profiler import profile_dataset
    from app.engine.validity import analyze_validity

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Transactions"
    ws.append(["transaction_id", "amount", "notes"])
    for i in range(19):
        # 2 blank notes cells (>5% missingness) + valid numeric amounts
        note_val = None if i < 2 else f"ok-{i}"
        ws.append([f"TXN-{i:04d}", round(100.0 + i * 5.25, 2), note_val])
    # 20th row has literal "N/A" in numeric amount column and empty string "" in notes
    ws.append(["TXN-0019", "N/A", ""])

    bio = io.BytesIO()
    wb.save(bio)
    xlsx_bytes = bio.getvalue()

    df, sheet, _ = parse_excel(xlsx_bytes, "xlsx")
    assert sheet == "Transactions"
    assert len(df) == 20
    # Literal "N/A" must remain present in amount (not coerced to NaN)
    assert (df["amount"].astype(str) == "N/A").sum() == 1
    # Ordinary blank cells (None and "") must still be NaN
    assert int(df["notes"].isna().sum()) == 3

    _, profiles = profile_dataset(df)
    val_findings = analyze_validity(df, profiles)
    cmp_findings = analyze_completeness(df, profiles)
    all_fids = {f.id for f in val_findings + cmp_findings}

    assert "FND-VAL-TYPEDRIFT-amount" in all_fids
    assert "FND-CMP-SENT-TXT-amount" in all_fids
    assert "FND-CMP-MISS-notes" in all_fids


def test_parse_pdf_retains_single_row_continuation_page_and_rejects_header_only(monkeypatch):
    """Verify multi-page PDF retains a valid 1-row final continuation page while rejecting repeated-header-only pages and 1-row PDFs."""
    import pdfplumber

    page1_table = [
        ["reading_id", "device_id", "temperature_c", "voltage", "vibration_mm_s"],
        ["RD00001", "SENS-001", "22.5", "5.01", "0.41"],
        ["RD00002", "SENS-002", "23.1", "4.99", "0.38"],
    ]
    page2_repeated_header_only = [
        ["reading_id", "device_id", "temperature_c", "voltage", "vibration_mm_s"],
    ]
    page3_single_continuation_row = [
        ["RD00003", "SENS-003", "4.59", "5.122", "0.801"],
    ]

    class DummyPage:
        def __init__(self, tables):
            self._tables = tables
        def extract_tables(self):
            return self._tables

    class DummyMultiPagePDF:
        def __init__(self):
            self.pages = [
                DummyPage([page1_table]),
                DummyPage([page2_repeated_header_only]),
                DummyPage([page3_single_continuation_row]),
            ]
        def __enter__(self):
            return self
        def __exit__(self, *args):
            pass

    monkeypatch.setattr(pdfplumber, "open", lambda _: DummyMultiPagePDF())
    df, page_count, _ = parse_pdf(b"%PDF-1.4 dummy")
    assert page_count == 3
    assert len(df) == 3
    assert df["reading_id"].tolist() == ["RD00001", "RD00002", "RD00003"]
    assert pd.api.types.is_float_dtype(df["temperature_c"])

    # Verify a PDF containing only a 1-row header table is still safely rejected
    class DummyHeaderOnlyPDF:
        def __init__(self):
            self.pages = [DummyPage([page2_repeated_header_only])]
        def __enter__(self):
            return self
        def __exit__(self, *args):
            pass

    monkeypatch.setattr(pdfplumber, "open", lambda _: DummyHeaderOnlyPDF())
    with pytest.raises(ValueError, match="machine-readable tabular data"):
        parse_pdf(b"%PDF-1.4 dummy")
